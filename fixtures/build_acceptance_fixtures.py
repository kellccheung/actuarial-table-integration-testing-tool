"""
Build acceptance fixtures for T01–T12.

Run from project root:
    python fixtures/build_acceptance_fixtures.py

Creates fixtures/acceptance/TXX_*/ with Control.xlsx, ChangeRequests/,
Production_Tables/, and (where useful) expected/ or pre-built ChangeLogs.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
import polars as pl

# Allow running as `python fixtures/build_acceptance_fixtures.py`
_PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from prophet_table_tool.prophet_csv import ProphetTable, write_prophet_csv

ROOT = Path(__file__).resolve().parent / "acceptance"


def _table(n_keys: int, columns: list[str], rows: list[list[str]]) -> ProphetTable:
    if rows:
        data = pl.DataFrame(rows, schema=columns, orient="row").cast(pl.Utf8)
    else:
        data = pl.DataFrame({c: [] for c in columns}).cast(pl.Utf8)
    return ProphetTable(n_keys=n_keys, columns=columns, data=data)


def _write_csv(path: Path, n_keys: int, columns: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    write_prophet_csv(_table(n_keys, columns, rows), path)


def _write_control(
    path: Path,
    *,
    working_root: Path,
    mode: str,
    run_id: str,
    change_requests: list[tuple],
    column_renames: list[tuple] | None = None,
    key_count_approvals: list[tuple] | None = None,
) -> None:
    """
    change_requests tuples:
        (change_request_id, order, include, approved, description, notes)
    column_renames tuples:
        (change_request_id, table_name, old_column_name, new_column_name)
    key_count_approvals tuples:
        (change_request_id, table_name, approved)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Config"
    ws.append(["key", "value"])
    ws.append(["working_root", str(working_root.resolve())])
    ws.append(["mode", mode])
    ws.append(["production_tables_path", "Production_Tables"])
    ws.append(["output_path", "Output"])
    ws.append(["run_id", run_id])

    ws_cr = wb.create_sheet("ChangeRequests")
    ws_cr.append(
        ["change_request_id", "order", "include", "approved", "description", "notes"]
    )
    for row in change_requests:
        ws_cr.append(list(row))

    ws_rn = wb.create_sheet("ColumnRenames")
    ws_rn.append(
        ["change_request_id", "table_name", "old_column_name", "new_column_name"]
    )
    for row in column_renames or []:
        ws_rn.append(list(row))

    ws_kc = wb.create_sheet("KeyCountApprovals")
    ws_kc.append(["change_request_id", "table_name", "approved"])
    for row in key_count_approvals or []:
        ws_kc.append(list(row))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _reset_case(name: str) -> Path:
    case_dir = ROOT / name
    if case_dir.exists():
        shutil.rmtree(case_dir)
    (case_dir / "ChangeRequests").mkdir(parents=True)
    (case_dir / "Production_Tables").mkdir(parents=True)
    (case_dir / "Output").mkdir(parents=True)
    (case_dir / "expected").mkdir(parents=True)
    return case_dir


def _cr_dirs(case_dir: Path, cr_id: str) -> tuple[Path, Path]:
    before = case_dir / "ChangeRequests" / cr_id / "before"
    after = case_dir / "ChangeRequests" / cr_id / "after"
    before.mkdir(parents=True)
    after.mkdir(parents=True)
    return before, after


# ---------------------------------------------------------------------------
# Shared sample schemas
# ---------------------------------------------------------------------------

MORT_COLS = ["Age", "Duration", "Product", "Rate", "Loading"]
MORT_N = 4  # !4 → keys: marker + Age, Duration, Product

EXP_COLS = ["Product", "ExpenseType", "Amount"]
EXP_N = 3

MORT_PROD = [
    ["20", "1", "PROD_A", "0.0012", "1.05"],
    ["25", "1", "PROD_A", "0.0015", "1.05"],
    ["30", "1", "PROD_B", "0.0020", "1.10"],
]

EXP_PROD = [
    ["PROD_A", "ACQ", "100"],
    ["PROD_A", "MAINT", "50"],
    ["PROD_B", "ACQ", "120"],
]


# ---------------------------------------------------------------------------
# T01 – Two CRs, no overlapping cells
# ---------------------------------------------------------------------------

def build_t01() -> Path:
    d = _reset_case("T01_two_crs_no_overlap")
    b1, a1 = _cr_dirs(d, "CR_T01_Mortality")
    b2, a2 = _cr_dirs(d, "CR_T01_Expense")

    # CR1: update Rate for Age=20 only
    mort_after = [
        ["20", "1", "PROD_A", "0.0013", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)

    # CR2: update Amount for PROD_A|ACQ only
    exp_after = [
        ["PROD_A", "ACQ", "110"],
        ["PROD_A", "MAINT", "50"],
        ["PROD_B", "ACQ", "120"],
    ]
    _write_csv(b2 / "EXPENSE_TABLE.csv", EXP_N, EXP_COLS, EXP_PROD)
    _write_csv(a2 / "EXPENSE_TABLE.csv", EXP_N, EXP_COLS, exp_after)

    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(d / "Production_Tables" / "EXPENSE_TABLE.csv", EXP_N, EXP_COLS, EXP_PROD)

    # Expected after both applied
    _write_csv(d / "expected" / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)
    _write_csv(d / "expected" / "EXPENSE_TABLE.csv", EXP_N, EXP_COLS, exp_after)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T01_RUN",
        change_requests=[
            ("CR_T01_Mortality", 1, "Y", "Y", "Mortality rate tweak", ""),
            ("CR_T01_Expense", 2, "Y", "Y", "Expense acquisition update", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T02 – Same cell conflict
# ---------------------------------------------------------------------------

def build_t02() -> Path:
    d = _reset_case("T02_same_cell_conflict")
    b1, a1 = _cr_dirs(d, "CR_T02_A")
    b2, a2 = _cr_dirs(d, "CR_T02_B")

    after_a = [
        ["20", "1", "PROD_A", "0.0018", "1.05"],  # conflict cell
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    after_b = [
        ["20", "1", "PROD_A", "0.0019", "1.05"],  # same key+col, different value
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    for folder, rows in ((b1, MORT_PROD), (a1, after_a), (b2, MORT_PROD), (a2, after_b)):
        _write_csv(folder / "MORT_TABLE.csv", MORT_N, MORT_COLS, rows)

    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T02_RUN",
        change_requests=[
            ("CR_T02_A", 1, "Y", "Y", "First conflicting update", ""),
            ("CR_T02_B", 2, "Y", "Y", "Second conflicting update", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T03 – Table only in before/
# ---------------------------------------------------------------------------

def build_t03() -> Path:
    d = _reset_case("T03_before_only_table")
    b1, a1 = _cr_dirs(d, "CR_T03_BeforeOnly")

    _write_csv(b1 / "OBSOLETE_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    # after/ intentionally empty (no CSVs)
    # Also include a normal change so the CR is not entirely empty
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    mort_after = [
        ["20", "1", "PROD_A", "0.0013", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)

    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(
        d / "Production_Tables" / "OBSOLETE_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD
    )

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T03_RUN",
        change_requests=[
            ("CR_T03_BeforeOnly", 1, "Y", "Y", "Before-only obsolete table", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T04 – Table only in after/ (table_add)
# ---------------------------------------------------------------------------

def build_t04() -> Path:
    d = _reset_case("T04_after_only_table_add")
    b1, a1 = _cr_dirs(d, "CR_T04_TableAdd")

    new_cols = ["Band", "Factor"]
    new_rows = [
        ["LOW", "0.90"],
        ["MID", "1.00"],
        ["HIGH", "1.10"],
    ]
    # before/ empty of NEW_LOADING_TABLE; after/ has it
    _write_csv(a1 / "NEW_LOADING_TABLE.csv", 2, new_cols, new_rows)

    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    # Expected row order matches apply()'s deterministic sort by key_tuple
    expected_rows = sorted(new_rows, key=lambda r: r[0])
    _write_csv(d / "expected" / "NEW_LOADING_TABLE.csv", 2, new_cols, expected_rows)
    _write_csv(d / "expected" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T04_RUN",
        change_requests=[
            ("CR_T04_TableAdd", 1, "Y", "Y", "Add new loading table", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T05 – Column rename declared
# ---------------------------------------------------------------------------

def build_t05() -> Path:
    d = _reset_case("T05_column_rename_declared")
    b1, a1 = _cr_dirs(d, "CR_T05_Rename")

    before_cols = ["Age", "Duration", "Product", "OldRate", "Loading"]
    after_cols = ["Age", "Duration", "Product", "Rate", "Loading"]
    before_rows = [
        ["20", "1", "PROD_A", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
    ]
    after_rows = [
        ["20", "1", "PROD_A", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, before_cols, before_rows)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, after_cols, after_rows)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, before_cols, before_rows)
    _write_csv(d / "expected" / "MORT_TABLE.csv", MORT_N, after_cols, after_rows)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T05_RUN",
        change_requests=[
            ("CR_T05_Rename", 1, "Y", "Y", "Rename OldRate to Rate", ""),
        ],
        column_renames=[
            ("CR_T05_Rename", "MORT_TABLE", "OldRate", "Rate"),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T06 – Column rename NOT declared (Stage 2 hard stop)
# ---------------------------------------------------------------------------

def build_t06() -> Path:
    """
    Stage 1 Control includes the rename so ChangeLog gets column_rename rows.
    Stage 2 uses Control_stage2.xlsx with ColumnRenames cleared → hard stop.
    """
    d = _reset_case("T06_column_rename_undeclared")
    b1, a1 = _cr_dirs(d, "CR_T06_Rename")

    before_cols = ["Age", "Duration", "Product", "OldRate", "Loading"]
    after_cols = ["Age", "Duration", "Product", "Rate", "Loading"]
    before_rows = [
        ["20", "1", "PROD_A", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
    ]
    after_rows = [
        ["20", "1", "PROD_A", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, before_cols, before_rows)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, after_cols, after_rows)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, before_cols, before_rows)

    # Stage 1 control (has rename declaration)
    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T06_RUN",
        change_requests=[
            ("CR_T06_Rename", 1, "Y", "Y", "Rename without Stage2 declaration", ""),
        ],
        column_renames=[
            ("CR_T06_Rename", "MORT_TABLE", "OldRate", "Rate"),
        ],
    )
    # Stage 2 control (NO column renames declared)
    _write_control(
        d / "Control_stage2.xlsx",
        working_root=d,
        mode="apply",
        run_id="T06_STAGE2",
        change_requests=[
            ("CR_T06_Rename", 1, "Y", "Y", "Rename without Stage2 declaration", ""),
        ],
        column_renames=[],  # intentionally empty
    )
    return d


# ---------------------------------------------------------------------------
# T07 – Key-count change approved = Y
# ---------------------------------------------------------------------------

def build_t07() -> Path:
    d = _reset_case("T07_key_count_approved")
    b1, a1 = _cr_dirs(d, "CR_T07_KeyCount")

    # before: !4 (Age, Duration, Product)  → after: !5 (+ Sex as new key)
    # One-to-one row mapping on Age|Duration|Product so diffs stay unambiguous.
    before_cols = ["Age", "Duration", "Product", "Rate", "Loading"]
    after_cols = ["Age", "Duration", "Product", "Sex", "Rate", "Loading"]
    before_rows = [
        ["20", "1", "PROD_A", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
    ]
    after_rows = [
        ["20", "1", "PROD_A", "M", "0.0012", "1.05"],
        ["25", "1", "PROD_A", "M", "0.0015", "1.05"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", 4, before_cols, before_rows)
    _write_csv(a1 / "MORT_TABLE.csv", 5, after_cols, after_rows)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", 4, before_cols, before_rows)
    _write_csv(d / "expected" / "MORT_TABLE.csv", 5, after_cols, after_rows)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T07_RUN",
        change_requests=[
            ("CR_T07_KeyCount", 1, "Y", "Y", "Add Sex as key column", ""),
        ],
        key_count_approvals=[
            ("CR_T07_KeyCount", "MORT_TABLE", "Y"),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T08 – Key-count change approved = N
# ---------------------------------------------------------------------------

def build_t08() -> Path:
    d = _reset_case("T08_key_count_not_approved")
    b1, a1 = _cr_dirs(d, "CR_T08_KeyCount")

    before_cols = ["Age", "Duration", "Product", "Rate", "Loading"]
    after_cols = ["Age", "Duration", "Product", "Sex", "Rate", "Loading"]
    before_rows = [["20", "1", "PROD_A", "0.0012", "1.05"]]
    after_rows = [["20", "1", "PROD_A", "M", "0.0012", "1.05"]]
    _write_csv(b1 / "MORT_TABLE.csv", 4, before_cols, before_rows)
    _write_csv(a1 / "MORT_TABLE.csv", 5, after_cols, after_rows)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", 4, before_cols, before_rows)

    # KeyCountApprovals=N blocks even when CR approved=Y
    # (is_key_count_approved checks KeyCountApprovals first).
    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T08_RUN",
        change_requests=[
            ("CR_T08_KeyCount", 1, "Y", "Y", "Unapproved key-count change", ""),
        ],
        key_count_approvals=[
            ("CR_T08_KeyCount", "MORT_TABLE", "N"),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T09 – validate_only (reuse T01-like data)
# ---------------------------------------------------------------------------

def build_t09() -> Path:
    d = _reset_case("T09_validate_only")
    b1, a1 = _cr_dirs(d, "CR_T09_Validate")

    mort_after = [
        ["20", "1", "PROD_A", "0.0013", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="validate_only",
        run_id="T09_RUN",
        change_requests=[
            ("CR_T09_Validate", 1, "Y", "Y", "Dry-run validation only", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T10 – Production old_value mismatch
# ---------------------------------------------------------------------------

def build_t10() -> Path:
    d = _reset_case("T10_old_value_mismatch")
    b1, a1 = _cr_dirs(d, "CR_T10_Mismatch")

    # before/after assume old Rate=0.0012
    mort_after = [
        ["20", "1", "PROD_A", "0.0013", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)

    # Production differs: Rate for 20|1|PROD_A is 0.0099 (not 0.0012)
    prod_mismatch = [
        ["20", "1", "PROD_A", "0.0099", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, prod_mismatch)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T10_RUN",
        change_requests=[
            ("CR_T10_Mismatch", 1, "Y", "Y", "Stale old_value vs production", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T11 – Idempotent double apply (same data as T01 subset)
# ---------------------------------------------------------------------------

def build_t11() -> Path:
    d = _reset_case("T11_idempotent_apply")
    b1, a1 = _cr_dirs(d, "CR_T11_Idempotent")

    mort_after = [
        ["20", "1", "PROD_A", "0.0013", "1.05"],
        ["25", "1", "PROD_A", "0.0015", "1.05"],
        ["30", "1", "PROD_B", "0.0020", "1.10"],
    ]
    _write_csv(b1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(a1 / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)
    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)
    _write_csv(d / "expected" / "MORT_TABLE.csv", MORT_N, MORT_COLS, mort_after)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T11_RUN",
        change_requests=[
            ("CR_T11_Idempotent", 1, "Y", "Y", "Idempotency check", ""),
        ],
    )
    return d


# ---------------------------------------------------------------------------
# T12 – Empty change request folder
# ---------------------------------------------------------------------------

def build_t12() -> Path:
    d = _reset_case("T12_empty_change_request")
    _cr_dirs(d, "CR_T12_Empty")  # before/ and after/ created, no CSVs

    _write_csv(d / "Production_Tables" / "MORT_TABLE.csv", MORT_N, MORT_COLS, MORT_PROD)

    _write_control(
        d / "Control.xlsx",
        working_root=d,
        mode="generate_changelog",
        run_id="T12_RUN",
        change_requests=[
            ("CR_T12_Empty", 1, "Y", "Y", "Empty CR folder", "no CSVs"),
        ],
    )
    return d


BUILDERS = [
    build_t01,
    build_t02,
    build_t03,
    build_t04,
    build_t05,
    build_t06,
    build_t07,
    build_t08,
    build_t09,
    build_t10,
    build_t11,
    build_t12,
]


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    print(f"Building fixtures under {ROOT}")
    for builder in BUILDERS:
        path = builder()
        print(f"  OK  {path.name}")
    print("Done.")


if __name__ == "__main__":
    main()
