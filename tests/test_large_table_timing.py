"""Large-table timing fixture for Stage 1 + Stage 2.

Builds a unique-key Prophet table in tmp_path (default 50_000 rows; override
with env ``PROPHET_TIMING_ROWS``), applies a sparse mix of updates / adds /
deletes, and prints wall times. Assertions check correctness plus a generous
time budget so a Python-loop regression fails CI.
"""

from __future__ import annotations

import os
import time
from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook, load_workbook

from prophet_table_tool.changelog import generate_change_log
from prophet_table_tool.integrate import integrate_changes
from prophet_table_tool.prophet_csv import ProphetTable, read_prophet_csv, write_prophet_csv

COLS = ["Age", "Duration", "Product", "Rate", "Loading"]
N_KEYS = 4
N_UPDATE = 200
N_DELETE = 50
N_ADD = 100
# Vectorized paths should finish well under this on a laptop; Python cell-loops would not.
BUDGET_S = 90.0


def _n_rows() -> int:
    raw = os.environ.get("PROPHET_TIMING_ROWS", "50000")
    n = int(raw)
    if n < N_UPDATE + N_DELETE + 1:
        raise ValueError(f"PROPHET_TIMING_ROWS must be > {N_UPDATE + N_DELETE}, got {n}")
    return n


def _large_table(n: int) -> ProphetTable:
    """Unique keys: Age = i % 100, Duration = i // 100, Product = PROD_A."""
    data = pl.DataFrame({"i": pl.int_range(0, n, eager=True)}).select(
        (pl.col("i") % 100).cast(pl.Utf8).alias("Age"),
        (pl.col("i") // 100).cast(pl.Utf8).alias("Duration"),
        pl.lit("PROD_A").alias("Product"),
        (pl.lit(0.001) + (pl.col("i") % 100).cast(pl.Float64) / 1e5)
        .cast(pl.Utf8)
        .alias("Rate"),
        pl.lit("1.05").alias("Loading"),
    )
    return ProphetTable(n_keys=N_KEYS, columns=COLS, data=data)


def _write_control(root: Path, run_id: str) -> Path:
    path = root / "Control.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Config"
    ws.append(["key", "value"])
    ws.append(["working_root", str(root.resolve())])
    ws.append(["mode", "generate_changelog"])
    ws.append(["production_tables_path", "Production_Tables"])
    ws.append(["output_path", "Output"])
    ws.append(["run_id", run_id])
    ws_cr = wb.create_sheet("ChangeRequests")
    ws_cr.append(
        ["change_request_id", "order", "include", "approved", "description", "notes"]
    )
    ws_cr.append(["CR_TIMING", 1, "Y", "Y", "large table timing", ""])
    wb.create_sheet("ColumnRenames").append(
        ["change_request_id", "table_name", "old_column_name", "new_column_name"]
    )
    wb.create_sheet("KeyCountApprovals").append(
        ["change_request_id", "table_name", "approved"]
    )
    wb.save(path)
    return path


@pytest.mark.timing
def test_large_table_stage1_stage2_timing(tmp_path: Path):
    n = _n_rows()
    cr = "CR_TIMING"
    prod_dir = tmp_path / "Production_Tables"
    before_dir = tmp_path / "ChangeRequests" / cr / "before"
    after_dir = tmp_path / "ChangeRequests" / cr / "after"
    prod_dir.mkdir(parents=True)
    before_dir.mkdir(parents=True)
    after_dir.mkdir(parents=True)

    t0 = time.perf_counter()
    prod = _large_table(n)
    write_prophet_csv(prod, prod_dir / "MORT_TABLE.csv")
    write_prophet_csv(prod, before_dir / "MORT_TABLE.csv")
    setup_s = time.perf_counter() - t0

    after_data = prod.data.with_columns(
        pl.when(pl.int_range(0, pl.len()) < N_UPDATE)
        .then(pl.lit("0.0099"))
        .otherwise(pl.col("Rate"))
        .alias("Rate")
    ).head(n - N_DELETE)
    add_i = pl.int_range(n, n + N_ADD, eager=True)
    added = pl.DataFrame({"i": add_i}).select(
        (pl.col("i") % 100).cast(pl.Utf8).alias("Age"),
        (pl.col("i") // 100).cast(pl.Utf8).alias("Duration"),
        pl.lit("PROD_A").alias("Product"),
        pl.lit("0.5").alias("Rate"),
        pl.lit("1.05").alias("Loading"),
    )
    after = ProphetTable(
        n_keys=N_KEYS,
        columns=COLS,
        data=pl.concat([after_data, added], how="vertical"),
    )
    write_prophet_csv(after, after_dir / "MORT_TABLE.csv")

    control = _write_control(tmp_path, "TIMING")

    t1 = time.perf_counter()
    clog = generate_change_log(control)
    stage1_s = time.perf_counter() - t1

    t2 = time.perf_counter()
    report = integrate_changes(control, clog, "apply")
    stage2_s = time.perf_counter() - t2
    total_s = setup_s + stage1_s + stage2_s

    print(
        f"\nlarge-table timing n={n} "
        f"setup={setup_s:.2f}s stage1={stage1_s:.2f}s stage2={stage2_s:.2f}s "
        f"total={total_s:.2f}s budget={BUDGET_S:.0f}s",
        flush=True,
    )

    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "SUCCESS"
    wb.close()

    out = read_prophet_csv(tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv")
    assert out.data.height == n - N_DELETE + N_ADD

    keyed = out.with_key_tuple()
    first_key = "0|0|PROD_A"
    updated = keyed.filter(pl.col("_key_str") == first_key)
    assert updated.height == 1
    assert updated["Rate"][0] == "0.0099"

    deleted_i = n - 1
    deleted_key = f"{deleted_i % 100}|{deleted_i // 100}|PROD_A"
    assert keyed.filter(pl.col("_key_str") == deleted_key).height == 0

    added_i = n
    added_key = f"{added_i % 100}|{added_i // 100}|PROD_A"
    added_row = keyed.filter(pl.col("_key_str") == added_key)
    assert added_row.height == 1
    assert added_row["Rate"][0] == "0.5"

    assert total_s < BUDGET_S, (
        f"large-table run took {total_s:.1f}s (budget {BUDGET_S:.0f}s); "
        "a Python-loop regression is likely"
    )
