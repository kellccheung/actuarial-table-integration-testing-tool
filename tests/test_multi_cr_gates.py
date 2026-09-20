"""Coverage for multi-CR conflicts, Control filters, deletes, and dry-run status."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from prophet_table_tool.changelog import generate_change_log
from prophet_table_tool.integrate import integrate_changes

from tests.test_sequential_apply import (
    COLS,
    N_KEYS,
    PROD,
    _cr_dirs,
    _mark_conflicts_resolved,
    _overlap_winner_rows,
    _report_messages,
    _stage2_leftovers,
    _write_control,
    _write_csv,
    _write_table,
)

PROD3 = [
    ["20", "1", "PROD_A", "0.0012", "1.05"],
    ["25", "1", "PROD_A", "0.0015", "1.05"],
    ["30", "1", "PROD_B", "0.0020", "1.10"],
]
NEW_ROW = ["99", "1", "PROD_A", "0.5", "1.05"]


def _summary_status(report: Path) -> str:
    wb = load_workbook(report, data_only=True)
    status = wb["Summary"]["B1"].value
    wb.close()
    return str(status)


def _conflict_rows(clog: Path) -> list[dict]:
    wb = load_workbook(clog, data_only=True)
    rows = list(wb["Conflicts"].iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in rows[0]]
    return [
        {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        for row in rows[1:]
        if row and any(c is not None for c in row)
    ]


def _set_cr_flags(
    control: Path, cr_id: str, *, include: str | None = None, approved: str | None = None
) -> None:
    wb = load_workbook(control)
    ws = wb["ChangeRequests"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    id_col = header.index("change_request_id") + 1
    include_col = header.index("include") + 1
    approved_col = header.index("approved") + 1
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, id_col).value == cr_id:
            if include is not None:
                ws.cell(row_idx, include_col, include)
            if approved is not None:
                ws.cell(row_idx, approved_col, approved)
    wb.save(control)
    wb.close()


def _assert_no_new_tables(root: Path) -> None:
    new_dir = root / "Output" / "New_Production_Tables"
    assert not new_dir.exists() or list(new_dir.glob("*.csv")) == []
    assert _stage2_leftovers(root) == []


def _write_mort(path: Path, rows: list[list[str]]) -> None:
    _write_csv(path / "MORT_TABLE.csv", rows)


def _setup_four_crs(tmp_path: Path) -> Path:
    """A/B overlap Age=20 Rate; C updates Age=25; D adds a row."""
    dirs = {cid: _cr_dirs(tmp_path, cid) for cid in ("CR_A", "CR_B", "CR_C", "CR_D")}
    _write_mort(tmp_path / "Production_Tables", PROD3)
    after_a = [PROD3[0][:], PROD3[1][:], PROD3[2][:]]
    after_a[0] = ["20", "1", "PROD_A", "0.0018", "1.05"]
    after_b = [PROD3[0][:], PROD3[1][:], PROD3[2][:]]
    after_b[0] = ["20", "1", "PROD_A", "0.0019", "1.05"]
    after_c = [PROD3[0][:], PROD3[1][:], PROD3[2][:]]
    after_c[1] = ["25", "1", "PROD_A", "0.0016", "1.05"]
    after_d = PROD3 + [NEW_ROW]
    for cid, rows_after in (
        ("CR_A", after_a),
        ("CR_B", after_b),
        ("CR_C", after_c),
        ("CR_D", after_d),
    ):
        before, after = dirs[cid]
        _write_mort(before, PROD3)
        _write_mort(after, rows_after)
    return _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Overlap first", ""),
            ("CR_B", 2, "Y", "Y", "Overlap second", ""),
            ("CR_C", 3, "Y", "Y", "Different cell", ""),
            ("CR_D", 4, "Y", "Y", "Add row", ""),
        ],
        "MULTI_FOUR",
    )


def test_same_table_different_cells_three_crs_no_conflict(tmp_path: Path):
    """A updates Age=20, B updates Age=25, C adds a row — no overlap, all apply."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    b3, a3 = _cr_dirs(tmp_path, "CR_C")
    after_a = [PROD3[0][:], PROD3[1][:], PROD3[2][:]]
    after_a[0] = ["20", "1", "PROD_A", "0.0013", "1.05"]
    after_b = [PROD3[0][:], PROD3[1][:], PROD3[2][:]]
    after_b[1] = ["25", "1", "PROD_A", "0.0016", "1.05"]
    after_c = PROD3 + [NEW_ROW]
    _write_mort(tmp_path / "Production_Tables", PROD3)
    _write_mort(b1, PROD3)
    _write_mort(a1, after_a)
    _write_mort(b2, PROD3)
    _write_mort(a2, after_b)
    _write_mort(b3, PROD3)
    _write_mort(a3, after_c)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Age 20", ""),
            ("CR_B", 2, "Y", "Y", "Age 25", ""),
            ("CR_C", 3, "Y", "Y", "Add 99", ""),
        ],
        "MULTI_NO_OVERLAP",
    )
    clog = generate_change_log(control)
    conflicts = _conflict_rows(clog)
    assert not any(r.get("conflict_type") for r in conflicts)

    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0013,1.05" in text
    assert "*,25,1,PROD_A,0.0016,1.05" in text
    assert "*,30,1,PROD_B,0.0020,1.10" in text
    assert "*,99,1,PROD_A,0.5,1.05" in text


def test_four_crs_unresolved_overlap_apply_fails(tmp_path: Path):
    control = _setup_four_crs(tmp_path)
    clog = generate_change_log(control)
    overlaps = [c for c in _conflict_rows(clog) if c.get("conflict_type") == "cell_overlap"]
    assert len(overlaps) == 1
    ids = {x.strip() for x in str(overlaps[0].get("change_request_ids") or "").split(",")}
    assert ids == {"CR_A", "CR_B"}

    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "FAILED"
    msgs = _report_messages(report)
    assert any("Unresolved conflict" in str(m.get("message") or "") for m in msgs)
    _assert_no_new_tables(tmp_path)


def test_four_crs_validate_only_unresolved_fails_and_writes_nothing(tmp_path: Path):
    control = _setup_four_crs(tmp_path)
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "validate_only")
    assert _summary_status(report) == "FAILED"
    msgs = _report_messages(report)
    assert any("Unresolved conflict" in str(m.get("message") or "") for m in msgs)
    _assert_no_new_tables(tmp_path)


def test_four_crs_resolved_later_overlap_wins_others_apply(tmp_path: Path):
    control = _setup_four_crs(tmp_path)
    clog = generate_change_log(control)
    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0019,1.05" in text
    assert "0.0018" not in text
    assert "*,25,1,PROD_A,0.0016,1.05" in text
    assert "*,30,1,PROD_B,0.0020,1.10" in text
    assert "*,99,1,PROD_A,0.5,1.05" in text
    assert _stage2_leftovers(tmp_path) == []


def test_three_crs_same_cell_resolved_last_writer_wins(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    b3, a3 = _cr_dirs(tmp_path, "CR_C")
    after_a = [["20", "1", "PROD_A", "0.0017", "1.05"], PROD[1][:]]
    after_b = [["20", "1", "PROD_A", "0.0018", "1.05"], PROD[1][:]]
    after_c = [["20", "1", "PROD_A", "0.0019", "1.05"], PROD[1][:]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    for before, after, rows in ((b1, a1, after_a), (b2, a2, after_b), (b3, a3, after_c)):
        _write_mort(before, PROD)
        _write_mort(after, rows)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "First", ""),
            ("CR_B", 2, "Y", "Y", "Second", ""),
            ("CR_C", 3, "Y", "Y", "Third", ""),
        ],
        "MULTI_THREE_WAY",
    )
    clog = generate_change_log(control)
    overlaps = [c for c in _conflict_rows(clog) if c.get("conflict_type") == "cell_overlap"]
    assert len(overlaps) == 1
    ids = {x.strip() for x in str(overlaps[0].get("change_request_ids") or "").split(",")}
    assert ids == {"CR_A", "CR_B", "CR_C"}

    report_fail = integrate_changes(control, clog, "apply")
    assert _summary_status(report_fail) == "FAILED"
    blocked = _overlap_winner_rows(report_fail)
    assert len(blocked) == 1
    assert blocked[0]["winner_change_request_id"] == "CR_C"
    assert blocked[0]["outcome"] == "blocked_unresolved"
    assert blocked[0]["superseded_change_request_ids"] == "CR_A, CR_B"
    assert blocked[0]["superseded_new_values"] == "CR_A=0.0017, CR_B=0.0018"

    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0019,1.05" in text
    assert "0.0017" not in text
    assert "0.0018" not in text
    winners = _overlap_winner_rows(report)
    assert len(winners) == 1
    row = winners[0]
    assert row["winner_change_request_id"] == "CR_C"
    assert row["winner_order"] == 3
    assert row["winner_new_value"] == "0.0019"
    assert row["overlapping_change_request_ids"] == "CR_A, CR_B, CR_C"
    assert row["superseded_change_request_ids"] == "CR_A, CR_B"
    assert row["superseded_new_values"] == "CR_A=0.0017, CR_B=0.0018"
    assert row["outcome"] == "applied"


def test_structural_collision_apply_fails_until_resolved(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    cols_new = COLS + ["NewCol"]
    cols_extra = COLS + ["ExtraCol"]
    with_new = [r + ["x"] for r in PROD]
    with_extra = [r + ["y"] for r in PROD]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_table(a1 / "MORT_TABLE.csv", cols_new, N_KEYS, with_new)
    _write_csv(b2 / "MORT_TABLE.csv", PROD)
    _write_table(a2 / "MORT_TABLE.csv", cols_extra, N_KEYS, with_extra)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Add NewCol", ""),
            ("CR_B", 2, "Y", "Y", "Add ExtraCol", ""),
        ],
        "MULTI_STRUCT",
    )
    clog = generate_change_log(control)
    structs = [
        c for c in _conflict_rows(clog) if c.get("conflict_type") == "structural_collision"
    ]
    assert structs
    ids = {x.strip() for x in str(structs[0].get("change_request_ids") or "").split(",")}
    assert ids == {"CR_A", "CR_B"}

    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "FAILED"
    _assert_no_new_tables(tmp_path)

    _mark_conflicts_resolved(clog)
    report_ok = integrate_changes(control, clog, "apply")
    assert _summary_status(report_ok) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "NewCol" in text
    assert "ExtraCol" in text
    assert "*,20,1,PROD_A,0.0012,1.05,x,y" in text


def test_row_delete_apply(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, [PROD[0]])
    control = _write_control(
        tmp_path, [("CR_A", 1, "Y", "Y", "Delete age 25", "")], "MULTI_ROW_DEL"
    )
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0012,1.05" in text
    assert "*,25,1,PROD_A," not in text


def test_column_delete_apply(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    cols_no_loading = ["Age", "Duration", "Product", "Rate"]
    after_rows = [r[:4] for r in PROD]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_table(a1 / "MORT_TABLE.csv", cols_no_loading, N_KEYS, after_rows)
    control = _write_control(
        tmp_path, [("CR_A", 1, "Y", "Y", "Drop Loading", "")], "MULTI_COL_DEL"
    )
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    header = text.splitlines()[0]
    assert "Loading" not in header
    assert "*,20,1,PROD_A,0.0012" in text


def test_row_delete_then_value_update_resolved_still_fails(tmp_path: Path):
    """CR_A deletes a row; CR_B later updates it — key is gone after sequential apply."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    after_b = [PROD[0][:], ["25", "1", "PROD_A", "0.0016", "1.05"]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, [PROD[0]])
    _write_mort(b2, PROD)
    _write_mort(a2, after_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Delete 25", ""),
            ("CR_B", 2, "Y", "Y", "Update 25", ""),
        ],
        "MULTI_DEL_THEN_UPD",
    )
    clog = generate_change_log(control)
    overlaps = [c for c in _conflict_rows(clog) if c.get("conflict_type") == "cell_overlap"]
    assert overlaps
    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "FAILED"
    msgs = _report_messages(report)
    assert any(
        m.get("change_request_id") == "CR_B"
        and "Key not found" in str(m.get("message") or "")
        for m in msgs
    )
    _assert_no_new_tables(tmp_path)


def test_stage1_include_n_skips_conflicting_cr(tmp_path: Path):
    """include=N at Stage 1 drops that CR, so a same-cell pair is not a conflict."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    after_a = [["20", "1", "PROD_A", "0.0018", "1.05"], PROD[1][:]]
    after_b = [["20", "1", "PROD_A", "0.0019", "1.05"], PROD[1][:]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, after_a)
    _write_mort(b2, PROD)
    _write_mort(a2, after_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Keep", ""),
            ("CR_B", 2, "N", "Y", "Excluded from Stage 1", ""),
        ],
        "MULTI_INCLUDE_N_S1",
    )
    clog = generate_change_log(control)
    assert not any(r.get("conflict_type") for r in _conflict_rows(clog))
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0018,1.05" in text
    assert "0.0019" not in text


def test_exclude_one_of_two_conflicting_crs_still_blocks(tmp_path: Path):
    """After Stage 1, setting include=N on one listed CR is not enough."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    after_a = [["20", "1", "PROD_A", "0.0018", "1.05"], PROD[1][:]]
    after_b = [["20", "1", "PROD_A", "0.0019", "1.05"], PROD[1][:]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, after_a)
    _write_mort(b2, PROD)
    _write_mort(a2, after_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "First", ""),
            ("CR_B", 2, "Y", "Y", "Second", ""),
        ],
        "MULTI_EXCLUDE_ONE",
    )
    clog = generate_change_log(control)
    _set_cr_flags(control, "CR_B", include="N")
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "FAILED"
    _assert_no_new_tables(tmp_path)


def test_exclude_all_listed_crs_allows_apply(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    b3, a3 = _cr_dirs(tmp_path, "CR_C")
    after_a = [["20", "1", "PROD_A", "0.0018", "1.05"], PROD[1][:]]
    after_b = [["20", "1", "PROD_A", "0.0019", "1.05"], PROD[1][:]]
    after_c = [PROD[0][:], ["25", "1", "PROD_A", "0.0016", "1.05"]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, after_a)
    _write_mort(b2, PROD)
    _write_mort(a2, after_b)
    _write_mort(b3, PROD)
    _write_mort(a3, after_c)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Overlap A", ""),
            ("CR_B", 2, "Y", "Y", "Overlap B", ""),
            ("CR_C", 3, "Y", "Y", "Independent", ""),
        ],
        "MULTI_EXCLUDE_ALL",
    )
    clog = generate_change_log(control)
    _set_cr_flags(control, "CR_A", include="N")
    _set_cr_flags(control, "CR_B", include="N")
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "SUCCESS"
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,25,1,PROD_A,0.0016,1.05" in text
    assert "0.0018" not in text
    assert "0.0019" not in text


def test_approved_n_on_one_conflicting_cr_still_blocks(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    after_a = [["20", "1", "PROD_A", "0.0018", "1.05"], PROD[1][:]]
    after_b = [["20", "1", "PROD_A", "0.0019", "1.05"], PROD[1][:]]
    _write_mort(tmp_path / "Production_Tables", PROD)
    _write_mort(b1, PROD)
    _write_mort(a1, after_a)
    _write_mort(b2, PROD)
    _write_mort(a2, after_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "First", ""),
            ("CR_B", 2, "Y", "Y", "Second", ""),
        ],
        "MULTI_APPROVED_N",
    )
    clog = generate_change_log(control)
    _set_cr_flags(control, "CR_B", approved="N")
    report = integrate_changes(control, clog, "apply")
    assert _summary_status(report) == "FAILED"
    _assert_no_new_tables(tmp_path)
