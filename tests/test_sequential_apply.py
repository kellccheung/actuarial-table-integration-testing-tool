"""Sequential Stage 2 validation: row_add then value_update, duplicate keys."""

from __future__ import annotations

from pathlib import Path

import polars as pl
from openpyxl import Workbook, load_workbook

from prophet_table_tool.changelog import generate_change_log
from prophet_table_tool.integrate import integrate_changes
from prophet_table_tool.prophet_csv import ProphetTable, write_prophet_csv

COLS = ["Age", "Duration", "Product", "Rate", "Loading"]
N_KEYS = 4
PROD = [
    ["20", "1", "PROD_A", "0.0012", "1.05"],
    ["25", "1", "PROD_A", "0.0015", "1.05"],
]
NEW_ROW = ["99", "1", "PROD_A", "0.5", "1.05"]
NEW_ROW_EDITED = ["99", "1", "PROD_A", "0.8", "1.05"]


def _write_table(path: Path, columns: list[str], n_keys: int, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = pl.DataFrame(rows, schema=columns, orient="row").cast(pl.Utf8)
    write_prophet_csv(ProphetTable(n_keys=n_keys, columns=columns, data=data), path)


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    _write_table(path, COLS, N_KEYS, rows)


def _write_control(root: Path, change_requests: list[tuple], run_id: str) -> Path:
    path = root / "Control.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Config"
    ws.append(["key", "value"])
    ws.append(["working_root", str(root.resolve())])
    ws.append(["mode", "generate_changelog"])
    ws.append(["production_tables_path", "Production_Tables"])
    ws.append(["output_path", "Output"])
    ws.append(["run_id", run_id])

    ws_cr = wb.create_sheet("ChangeRequests")
    ws_cr.append(
        ["change_request_id", "order", "include", "approved", "description", "notes"]
    )
    for row in change_requests:
        ws_cr.append(list(row))

    wb.create_sheet("ColumnRenames").append(
        ["change_request_id", "table_name", "old_column_name", "new_column_name"]
    )
    wb.create_sheet("KeyCountApprovals").append(
        ["change_request_id", "table_name", "approved"]
    )
    wb.save(path)
    return path


def _cr_dirs(root: Path, cr_id: str) -> tuple[Path, Path]:
    before = root / "ChangeRequests" / cr_id / "before"
    after = root / "ChangeRequests" / cr_id / "after"
    before.mkdir(parents=True)
    after.mkdir(parents=True)
    return before, after


def _mark_conflicts_resolved(clog: Path) -> None:
    wb = load_workbook(clog)
    ws = wb["Conflicts"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    resolved_col = header.index("resolved") + 1
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, 1).value:
            ws.cell(row_idx, resolved_col, "Y")
    wb.save(clog)
    wb.close()


def _report_messages(report: Path) -> list[dict]:
    wb = load_workbook(report, data_only=True)
    ws = wb["Validation_Report"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in rows[0]]
    return [
        {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        for row in rows[1:]
        if row and any(c is not None for c in row)
    ]


def _setup_add_then_edit(tmp_path: Path) -> Path:
    """CR_A adds a row; CR_B's before includes that row and edits Rate."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    prod_plus = PROD + [NEW_ROW]
    prod_plus_edit = PROD + [NEW_ROW_EDITED]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_csv(a1 / "MORT_TABLE.csv", prod_plus)
    _write_csv(b2 / "MORT_TABLE.csv", prod_plus)
    _write_csv(a2 / "MORT_TABLE.csv", prod_plus_edit)
    return _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Add row", ""),
            ("CR_B", 2, "Y", "Y", "Edit new row", ""),
        ],
        "SEQ_ADD_EDIT",
    )


def test_row_add_then_value_update_apply_after_resolved(tmp_path: Path):
    control = _setup_add_then_edit(tmp_path)
    clog = generate_change_log(control)
    wb = load_workbook(clog, data_only=True)
    conflicts = list(wb["Conflicts"].iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in conflicts[0]]
    type_idx = header.index("conflict_type")
    notes_idx = header.index("notes")
    overlap_notes = [
        row[notes_idx]
        for row in conflicts[1:]
        if row and row[type_idx] == "cell_overlap"
    ]
    assert overlap_notes
    assert any("row_add and value_update" in str(n) for n in overlap_notes)

    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "SUCCESS"
    wb.close()

    out = tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv"
    text = out.read_text(encoding="utf-8")
    assert "*,99,1,PROD_A,0.8,1.05" in text
    assert text.count("*,99,1,PROD_A,") == 1


def test_row_add_then_value_update_validate_only_no_files(tmp_path: Path):
    control = _setup_add_then_edit(tmp_path)
    clog = generate_change_log(control)
    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "validate_only")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "DRY_RUN_SUCCESS"
    wb.close()
    new_dir = tmp_path / "Output" / "New_Production_Tables"
    assert not new_dir.exists() or list(new_dir.glob("*.csv")) == []


def test_row_add_then_value_update_unresolved_apply_fails(tmp_path: Path):
    control = _setup_add_then_edit(tmp_path)
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "FAILED"
    wb.close()
    new_dir = tmp_path / "Output" / "New_Production_Tables"
    assert not new_dir.exists() or list(new_dir.glob("*.csv")) == []


def test_duplicate_row_add_fails_key_already_exists(tmp_path: Path):
    """Both CRs row_add the same key from production; second add is rejected."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    added = PROD + [NEW_ROW]
    added_b = PROD + [NEW_ROW_EDITED]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_csv(a1 / "MORT_TABLE.csv", added)
    _write_csv(b2 / "MORT_TABLE.csv", PROD)
    _write_csv(a2 / "MORT_TABLE.csv", added_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Add row A", ""),
            ("CR_B", 2, "Y", "Y", "Add same row B", ""),
        ],
        "SEQ_DUP_ADD",
    )
    clog = generate_change_log(control)
    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "FAILED"
    wb.close()
    msgs = _report_messages(report)
    assert any(
        m.get("change_type") == "row_add"
        and "already exists" in str(m.get("message") or "")
        for m in msgs
    )
    new_dir = tmp_path / "Output" / "New_Production_Tables"
    assert not new_dir.exists() or list(new_dir.glob("*.csv")) == []


def test_row_add_existing_production_key_fails(tmp_path: Path):
    """CR before omits a production row so Stage 1 emits row_add; Stage 2 rejects it."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", [PROD[0]])
    _write_csv(a1 / "MORT_TABLE.csv", PROD)
    control = _write_control(
        tmp_path,
        [("CR_A", 1, "Y", "Y", "Re-add existing key", "")],
        "SEQ_EXISTING_ADD",
    )
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "FAILED"
    wb.close()
    msgs = _report_messages(report)
    assert any("already exists" in str(m.get("message") or "") for m in msgs)


def test_value_update_missing_key_still_fails(tmp_path: Path):
    b1, a1 = _cr_dirs(tmp_path, "CR_B")
    prod_plus = PROD + [NEW_ROW]
    prod_plus_edit = PROD + [NEW_ROW_EDITED]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", prod_plus)
    _write_csv(a1 / "MORT_TABLE.csv", prod_plus_edit)
    control = _write_control(
        tmp_path,
        [("CR_B", 1, "Y", "Y", "Edit row that is not in production", "")],
        "SEQ_MISSING_UPDATE",
    )
    clog = generate_change_log(control)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "FAILED"
    wb.close()
    msgs = _report_messages(report)
    assert any("Key not found in current table" in str(m.get("message") or "") for m in msgs)


COLS_WITH_NEW = COLS + ["NewCol"]


def _setup_column_add_then_edit(tmp_path: Path) -> Path:
    """CR_A adds NewCol; CR_B's before includes NewCol and changes a cell."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    with_col = [r + ["x"] for r in PROD]
    with_col_edit = [PROD[0] + ["y"], PROD[1] + ["x"]]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_table(a1 / "MORT_TABLE.csv", COLS_WITH_NEW, N_KEYS, with_col)
    _write_table(b2 / "MORT_TABLE.csv", COLS_WITH_NEW, N_KEYS, with_col)
    _write_table(a2 / "MORT_TABLE.csv", COLS_WITH_NEW, N_KEYS, with_col_edit)
    return _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "Add column", ""),
            ("CR_B", 2, "Y", "Y", "Edit new column", ""),
        ],
        "SEQ_COL_ADD_EDIT",
    )


def test_column_add_then_value_update_apply_after_resolved(tmp_path: Path):
    control = _setup_column_add_then_edit(tmp_path)
    clog = generate_change_log(control)
    wb = load_workbook(clog, data_only=True)
    conflicts = list(wb["Conflicts"].iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in conflicts[0]]
    type_idx = header.index("conflict_type")
    notes_idx = header.index("notes")
    overlap_notes = [
        row[notes_idx]
        for row in conflicts[1:]
        if row and row[type_idx] == "cell_overlap"
    ]
    assert overlap_notes
    assert any("column_add and value_update" in str(n) for n in overlap_notes)

    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "SUCCESS"
    wb.close()

    out = tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv"
    text = out.read_text(encoding="utf-8")
    assert "*,20,1,PROD_A,0.0012,1.05,y" in text
    assert "*,25,1,PROD_A,0.0015,1.05,x" in text


def test_differing_value_update_resolved_later_cr_wins(tmp_path: Path):
    """Differing cell_overlap: resolved=Y applies Control order; later CR wins."""
    b1, a1 = _cr_dirs(tmp_path, "CR_A")
    b2, a2 = _cr_dirs(tmp_path, "CR_B")
    after_a = [PROD[0][:], PROD[1][:]]
    after_a[0] = ["20", "1", "PROD_A", "0.0018", "1.05"]
    after_b = [PROD[0][:], PROD[1][:]]
    after_b[0] = ["20", "1", "PROD_A", "0.0019", "1.05"]
    _write_csv(tmp_path / "Production_Tables" / "MORT_TABLE.csv", PROD)
    _write_csv(b1 / "MORT_TABLE.csv", PROD)
    _write_csv(a1 / "MORT_TABLE.csv", after_a)
    _write_csv(b2 / "MORT_TABLE.csv", PROD)
    _write_csv(a2 / "MORT_TABLE.csv", after_b)
    control = _write_control(
        tmp_path,
        [
            ("CR_A", 1, "Y", "Y", "First rate", ""),
            ("CR_B", 2, "Y", "Y", "Second rate", ""),
        ],
        "SEQ_LAST_WRITER",
    )
    clog = generate_change_log(control)
    _mark_conflicts_resolved(clog)
    report = integrate_changes(control, clog, "apply")
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B1"].value == "SUCCESS"
    wb.close()
    text = (tmp_path / "Output" / "New_Production_Tables" / "MORT_TABLE.csv").read_text(
        encoding="utf-8"
    )
    assert "*,20,1,PROD_A,0.0019,1.05" in text
    assert "0.0018" not in text
